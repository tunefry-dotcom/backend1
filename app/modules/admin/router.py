"""Admin router — internal user management and submission review endpoints.

All routes require X-Admin-Secret header matching settings.admin_secret.
"""

from __future__ import annotations

import asyncio
import logging
import re
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Annotated, Any, List, Literal, Optional
from uuid import uuid4

import openpyxl
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Query, UploadFile, status
from pydantic import BaseModel, Field
from starlette.concurrency import run_in_threadpool

from app.core.config import settings
from app.core.email import send_email, submission_review_email_html
from app.core.r2_client import delete_keys, presign_get, upload_bytes
from app.core.supabase_client import get_service_client
from app.modules.billing.plans import Plan, get_spec
from app.modules.billing.service import assign_plan
from app.modules.blog import ai as blog_ai
from app.modules.blog import service as blog_service
from app.modules.blog.schemas import (
    AdminFlagsRequest,
    AdminPost,
    AdminReviewRequest,
    AdminTunefryPublishRequest,
    RewriteRequest,
    RewriteResponse,
)
from app.modules.blog.service import ReviewError
from app.modules.home import service as home_service
from app.modules.home.schemas import HomeContent
from app.modules.earnings.service import get_balance, recompute_balance
from app.modules.profile import service as profile_service
from app.modules.referrals.service import credit_referral
from migration.platform_map import normalize_platform

_log = logging.getLogger(__name__)

router = APIRouter(prefix="/admin", tags=["admin"])

_PLAN_NAMES: dict[str, str] = {
    "free": "Free",
    "single-song": "Single Song",
    "starter": "Starter",
    "single-artist": "Single Artist",
    "double-artist": "Double Artist",
    "label": "Label",
    # legacy underscore variants (kept for backwards compat)
    "single_song": "Single Song",
    "single_artist": "Single Artist",
    "double_artist": "Double Artist",
}

_PLAN_PRICES_INR: dict[str, int] = {
    "single-song": 299,
    "starter": 999,
    "single-artist": 1599,
    "double-artist": 2999,
    "label": 6999,
    "single_song": 299,
    "single_artist": 1599,
    "double_artist": 2999,
}


def _require_admin(x_admin_secret: Annotated[str, Header()] = "") -> None:
    if not settings.admin_secret or x_admin_secret != settings.admin_secret:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")


def _fmt(val: Any) -> str | None:
    if val is None:
        return None
    return val.isoformat() if hasattr(val, "isoformat") else str(val)


def _fetch_all_users(svc: Any) -> list:
    """Paginate through auth.admin.list_users() to get every user."""
    users: list = []
    page = 1
    per_page = 1000
    while True:
        batch = svc.auth.admin.list_users(page=page, per_page=per_page)
        if not batch:
            break
        users.extend(batch)
        if len(batch) < per_page:
            break
        page += 1
    return users


def _fetch_all_rows(svc: Any, table: str, columns: str) -> list:
    """Paginate through a PostgREST table to bypass the default 1 000-row cap.

    PostgREST returns at most 1 000 rows by default.  With 2 300+ users each
    having a subscriptions/profiles row, any plain .execute() silently drops
    the tail — meaning those users always appear as free/blank in the admin
    panel.  This function pages through with .range() until all rows are in.
    """
    rows: list = []
    page_size = 1000
    start = 0
    while True:
        resp = (
            svc.table(table)
            .select(columns)
            .range(start, start + page_size - 1)
            .execute()
        )
        batch = resp.data or []
        rows.extend(batch)
        if len(batch) < page_size:
            break
        start += page_size
    return rows


@router.get("/users", dependencies=[Depends(_require_admin)])
async def list_users(q: str = Query(default="")) -> dict:
    """Return all users with their plan and subscription status.

    Optional ?q= filters by email or full_name (case-insensitive).
    """
    try:
        svc = get_service_client()
        raw_users = _fetch_all_users(svc)
        all_subs = _fetch_all_rows(
            svc, "subscriptions", "user_id,plan,status,expires_at,started_at"
        )
        all_profiles = _fetch_all_rows(
            svc, "profiles",
            "user_id,full_name,artist_name,phone,spotify_url,apple_music_url,instagram,youtube_url,city,state,bio,gender,date_of_birth,custom_label_name",
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch users: {exc}",
        ) from exc

    sub_map: dict[str, dict] = {row["user_id"]: row for row in all_subs}
    profile_map: dict[str, dict] = {row["user_id"]: row for row in all_profiles}

    users = []
    for u in raw_users:
        uid = str(getattr(u, "id", "") or "")
        email = getattr(u, "email", "") or ""
        app_meta: dict = getattr(u, "app_metadata", None) or {}
        user_meta: dict = getattr(u, "user_metadata", None) or {}
        full_name: str = user_meta.get("full_name", "") or ""
        artist_name: str = user_meta.get("artist_name", "") or ""
        phone: str = user_meta.get("phone", "") or ""
        sub = sub_map.get(uid, {})
        plan: str = sub.get("plan") or app_meta.get("plan", "free") or "free"
        prof = profile_map.get(uid, {})
        # Profiles table is authoritative; fall back to user_meta for legacy users
        # whose profiles row predates admin edits writing these fields there.
        full_name = prof.get("full_name") or full_name
        artist_name = prof.get("artist_name") or artist_name
        phone = prof.get("phone") or phone

        users.append(
            {
                "id": uid,
                "email": email,
                "full_name": full_name,
                "artist_name": artist_name,
                "phone": phone,
                "plan": plan,
                "plan_name": _PLAN_NAMES.get(plan, plan.replace("_", " ").title()),
                "status": sub.get("status", "active"),
                "expires_at": sub.get("expires_at"),
                "created_at": _fmt(getattr(u, "created_at", None)),
                "last_sign_in_at": _fmt(getattr(u, "last_sign_in_at", None)),
                "spotify_url": prof.get("spotify_url") or "",
                "apple_music_url": prof.get("apple_music_url") or "",
                "instagram": prof.get("instagram") or "",
                "youtube_url": prof.get("youtube_url") or "",
                "city": prof.get("city") or "",
                "state": prof.get("state") or "",
                "bio": prof.get("bio") or "",
                "gender": prof.get("gender") or "",
                "date_of_birth": prof.get("date_of_birth") or "",
                "custom_label_name": prof.get("custom_label_name") or "",
            }
        )

    if q:
        q_lower = q.lower()
        users = [
            u for u in users
            if q_lower in u["email"].lower() or q_lower in u["full_name"].lower()
        ]

    return {"users": users, "total": len(users)}


# ---------------------------------------------------------------------------
# Submission review
# ---------------------------------------------------------------------------

_CATEGORY_TYPES: dict[str, list[str]] = {
    "new-songs":        ["new_song"],
    "transfer-songs":   ["transfer_song"],
    "new-albums":       ["new_album"],
    "transfer-albums":  ["transfer_album"],
    "profile-mismatch": ["profile_mismatch"],
    "claim-removal":    ["claim_removal"],
    "insta-link":       ["insta_link"],
}


class ReviewBody(BaseModel):
    status: str        # "approved" | "declined"
    admin_note: str = ""


class BroadcastBody(BaseModel):
    title: str
    body: str = ""


def _submission_title(data: dict) -> str:
    """Best-effort display title for a submission (mirrors the frontend subTitle)."""
    for key in ("song_title", "album_name", "section_name", "song_name", "instagram_url"):
        val = data.get(key)
        if val:
            return str(val)
    return ""


def _submission_keys(data: dict) -> set[str]:
    """All R2 object keys referenced by a submission's data payload."""
    keys = {data.get("cover_art_key"), data.get("audio_key")}
    for song in (data.get("songs") or []):
        keys.add(song.get("audio_key"))
    return {k for k in keys if k}


@router.get("/submissions/{category}", dependencies=[Depends(_require_admin)])
async def list_submissions(
    category: str,
    page: int = Query(default=1, ge=1),
    per_page: int = Query(default=10, ge=1, le=50),
    q: str = Query(default=""),
    plan: str = Query(default=""),
) -> dict:
    """Paginated submissions for a category.

    Sorted by created_at DESC only — status never affects ordering, so a
    card's position stays stable across approve/decline and reload.
    """
    types = _CATEGORY_TYPES.get(category)
    if not types:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unknown category")

    svc = get_service_client()

    try:
        resp = (
            svc.table("submissions")
            .select("*")
            .in_("submission_type", types)
            .order("created_at", desc=True)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch submissions: {exc}",
        ) from exc

    all_items = resp.data or []

    # Overwrite the stored user_plan snapshot with the user's *current* plan,
    # joined live from subscriptions (same source as the All Users tab). The
    # stored value is captured from the JWT at submission time and is stale
    # (or free) after upgrades or if the access-token hook isn't stamping it.
    try:
        raw_users = _fetch_all_users(svc)
        all_subs = _fetch_all_rows(
            svc, "subscriptions", "user_id,plan,status,expires_at,started_at"
        )
        sub_map: dict[str, dict] = {row["user_id"]: row for row in all_subs}
        email_plan: dict[str, str] = {}
        for u in raw_users:
            uid = str(getattr(u, "id", "") or "")
            email = (getattr(u, "email", "") or "").lower()
            if email:
                email_plan[email] = sub_map.get(uid, {}).get("plan") or "free"
        for item in all_items:
            key = (item.get("user_email") or "").lower()
            item["user_plan"] = email_plan.get(key, item.get("user_plan") or "free")
    except Exception:
        pass  # best-effort — fall back to the stored snapshot on any failure

    # Filter by the user's (live-joined) plan.
    if plan and plan.lower() != "all":
        all_items = [it for it in all_items if (it.get("user_plan") or "free") == plan]

    # Search by song/album title or user email (case-insensitive substring).
    query = q.strip().lower()
    if query:
        def _matches(it: dict) -> bool:
            title = _submission_title(it.get("data") or {}).lower()
            email = (it.get("user_email") or "").lower()
            return query in title or query in email
        all_items = [it for it in all_items if _matches(it)]

    total = len(all_items)
    offset = (page - 1) * per_page
    total_pages = max(1, -(-total // per_page))
    return {
        "submissions": all_items[offset: offset + per_page],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": total_pages,
    }


@router.patch("/submissions/{submission_id}", dependencies=[Depends(_require_admin)])
async def review_submission(submission_id: str, body: ReviewBody) -> dict:
    """Approve or decline a submission."""
    if body.status not in ("approved", "declined"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="status must be 'approved' or 'declined'",
        )

    svc = get_service_client()
    try:
        resp = (
            svc.table("submissions")
            .update(
                {
                    "status": body.status,
                    "admin_note": body.admin_note,
                    "reviewed_at": datetime.now(timezone.utc).isoformat(),
                }
            )
            .eq("id", submission_id)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not update submission: {exc}",
        ) from exc

    if not resp.data:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Submission not found")

    submission = resp.data[0]

    # If approving a new-artist submission, add to the new-artist queue.
    if body.status == "approved":
        data: dict = submission.get("data") or {}
        if str(data.get("new_artist", "")).lower() == "true":
            main_artists = data.get("main_artists") or []
            if not main_artists:
                songs = data.get("songs") or []
                if songs:
                    main_artists = songs[0].get("main_artists") or []
            artist_name = (main_artists[0].get("name", "") if main_artists else "") or ""
            try:
                svc.table("new_artist_queue").insert({
                    "user_email": submission.get("user_email", ""),
                    "artist_name": artist_name,
                    "submission_id": submission_id,
                }).execute()
            except Exception:
                pass  # best-effort — don't block the approval

    # Fire-and-forget email to the artist (non-blocking — approval/decline is never held back).
    _user_email: str = submission.get("user_email") or ""
    if _user_email:
        _title = _submission_title(submission.get("data") or {}) or (
            submission.get("submission_type", "submission").replace("_", " ").title()
        )
        _subject = (
            "Your submission was approved 🎉"
            if body.status == "approved"
            else f"Submission update — {_title}"
        )
        try:
            asyncio.create_task(
                send_email(
                    to=_user_email,
                    subject=_subject,
                    html_body=submission_review_email_html(body.status, _title, body.admin_note),
                )
            )
        except Exception:
            _log.warning("Could not schedule submission review email", exc_info=True)

    return submission


@router.post("/notifications", dependencies=[Depends(_require_admin)])
async def create_notification(body: BroadcastBody) -> dict:
    """Push a broadcast announcement to all users' notification bell."""
    if not body.title.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="title is required")
    svc = get_service_client()
    try:
        resp = (
            svc.table("notifications")
            .insert({"title": body.title.strip(), "body": body.body.strip()})
            .execute()
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not create notification: {exc}",
        ) from exc
    row = resp.data[0] if resp.data else {}
    return {"ok": True, "id": row.get("id")}


# ---------------------------------------------------------------------------
# Withdrawal requests
# ---------------------------------------------------------------------------
def _to_decimal(v: Any) -> Decimal:
    try:
        return Decimal(str(v)) if v is not None else Decimal("0")
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _age_from_dob(dob: Any) -> int | None:
    if not dob:
        return None
    try:
        d = dob if isinstance(dob, date) else datetime.fromisoformat(str(dob)[:10]).date()
    except (ValueError, TypeError):
        return None
    today = date.today()
    return today.year - d.year - ((today.month, today.day) < (d.month, d.day))


def _parse_expires_at(value: str) -> str | None:
    """Convert admin-supplied date string to a UTC end-of-day ISO timestamp.

    ""           → None  (store NULL — plan never expires)
    "YYYY-MM-DD" → "YYYY-MM-DDT23:59:59+00:00"

    End-of-day storage ensures effective_plan()'s ``expires_at <= now()`` check
    does not expire the plan at 00:00:00 UTC on the chosen day — the plan stays
    active through the full calendar day.
    """
    if not value:
        return None
    try:
        d = date.fromisoformat(value)
    except ValueError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"expires_at must be YYYY-MM-DD or empty string, got: {value!r}",
        )
    eod = datetime(d.year, d.month, d.day, 23, 59, 59, tzinfo=timezone.utc)
    return eod.isoformat()


@router.get("/withdrawals", dependencies=[Depends(_require_admin)])
async def list_withdrawals() -> dict:
    """All withdrawal requests, pending first then newest — with the artist
    snapshot (plan / name / address / age) and payout details for payout."""
    svc = get_service_client()
    try:
        rows = _fetch_all_rows(
            svc, "withdrawal_requests",
            "id,user_id,user_email,amount,status,method,payout_details,"
            "snapshot,admin_note,requested_at,processed_at",
        )
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=f"Could not fetch withdrawals: {exc}") from exc
    # pending first; within each group, most-recently requested first
    rows.sort(key=lambda r: _fmt(r.get("requested_at")) or "", reverse=True)
    rows.sort(key=lambda r: r.get("status") != "pending")
    uids = list({r["user_id"] for r in rows if r.get("user_id")})
    # Backfill missing snapshot fields from live profile data (graceful — never blocks).
    try:
        if uids:
            profiles_raw = (
                svc.table("profiles")
                .select("id,phone,city,state,date_of_birth")
                .in_("id", uids)
                .execute()
                .data or []
            )
            prof_map = {p["id"]: p for p in profiles_raw}
            for row in rows:
                prof = prof_map.get(row.get("user_id"), {})
                snap = row.get("snapshot") or {}
                if not snap.get("phone"):
                    snap["phone"] = prof.get("phone")
                if not snap.get("city"):
                    snap["city"] = prof.get("city")
                if not snap.get("state"):
                    snap["state"] = prof.get("state")
                if snap.get("age") is None:
                    snap["age"] = _age_from_dob(prof.get("date_of_birth"))
                row["snapshot"] = snap
    except Exception as exc:
        _log.warning("Could not backfill profile fields for withdrawals list: %s", exc)
    # Live plan join (same source as /admin/users and /admin/submissions) — kept in its
    # own try/except so a profile-fetch failure above can't silently suppress this too.
    # snapshot["plan"] is frozen at request time and goes stale after upgrades.
    try:
        if uids:
            subs_raw = _fetch_all_rows(
                svc, "subscriptions", "user_id,plan,status,expires_at,started_at"
            )
            sub_map = {s["user_id"]: s for s in subs_raw}
            for row in rows:
                plan = sub_map.get(row.get("user_id"), {}).get("plan") or "free"
                try:
                    plan_enum = Plan(plan)
                except ValueError:
                    plan_enum = Plan("free")
                snap = row.get("snapshot") or {}
                snap["plan"] = plan_enum.value
                snap["plan_name"] = get_spec(plan_enum).name
                row["snapshot"] = snap
    except Exception as exc:
        _log.warning("Could not live-join plan for withdrawals list: %s", exc)
    total_pending = sum(_to_decimal(r["amount"]) for r in rows if r.get("status") == "pending")
    return {"requests": rows, "total_pending": float(total_pending)}


class WithdrawalReviewBody(BaseModel):
    status: Literal["paid", "declined"]
    admin_note: str = Field(min_length=1, max_length=500)


@router.patch("/withdrawals/{request_id}", dependencies=[Depends(_require_admin)])
async def review_withdrawal(request_id: str, body: WithdrawalReviewBody) -> dict:
    """Mark a request as paid or declined; both require a comment shown to the artist.

    Paid: balance was already zeroed at request time, so no balance change is
    needed here (the amount stays counted as withdrawn). Declined: the amount
    was reserved (zeroed out of available_balance) but never added to
    total_withdrawn, so recompute_balance's pending-only subtraction naturally
    releases it back once the row's status leaves "pending" — no manual
    arithmetic, and total_withdrawn is untouched.
    """
    svc = get_service_client()
    try:
        existing_res = (svc.table("withdrawal_requests")
                        .select("user_email").eq("id", request_id).limit(1).execute())
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=f"Could not fetch withdrawal: {exc}") from exc
    existing = (existing_res.data or [None])[0]
    if not existing:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Request not found")

    try:
        resp = (svc.table("withdrawal_requests").update({
            "status": body.status,
            "admin_note": body.admin_note,
            "processed_at": datetime.now(timezone.utc).isoformat(),
        }).eq("id", request_id).execute())
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=f"Could not update withdrawal: {exc}") from exc
    if not resp.data:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Request not found")

    balance = None
    if body.status == "declined":
        balance = recompute_balance((existing.get("user_email") or "").lower())
    return {"ok": True, "request": resp.data[0], "balance": balance}


@router.delete("/withdrawals/{request_id}", dependencies=[Depends(_require_admin)])
async def delete_withdrawal(request_id: str) -> dict:
    """Delete a request. If it is still pending, credit the reserved amount back
    to the artist's available_balance so no earnings are lost (the request had
    zeroed it). Paid or declined requests never credit back here — paid stays
    counted as withdrawn, and declined already had its amount restored via
    recompute_balance in review_withdrawal (crediting again would double-count)."""
    svc = get_service_client()
    try:
        res = (svc.table("withdrawal_requests")
               .select("user_email,amount,status").eq("id", request_id).limit(1).execute())
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=f"Could not fetch withdrawal: {exc}") from exc
    row = (res.data or [None])[0]
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Request not found")

    if row.get("status") == "pending":
        email = (row.get("user_email") or "").lower()
        bal = (svc.table("artist_balances").select("available_balance")
               .eq("user_email", email).limit(1).execute())
        current = _to_decimal((bal.data or [{}])[0].get("available_balance"))
        new_balance = current + _to_decimal(row.get("amount"))
        svc.table("artist_balances").update({
            "available_balance": str(new_balance),
            "last_updated": datetime.now(timezone.utc).isoformat(),
        }).eq("user_email", email).execute()

    svc.table("withdrawal_requests").delete().eq("id", request_id).execute()
    return {"ok": True, "deleted": request_id}


class DeleteBody(BaseModel):
    ids: list[str]


@router.delete("/submissions", dependencies=[Depends(_require_admin)])
async def delete_submissions(body: DeleteBody) -> dict:
    """Delete one or more submissions (DB rows) and their orphaned R2 files.

    R2 object keys are derived from {artist}/{release} and are NOT unique per
    submission, so a key is only removed from R2 if no *surviving* submission
    still references it (otherwise deleting one row would break another's files).
    """
    ids = [i for i in body.ids if i]
    if not ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No ids provided")

    svc = get_service_client()

    # 1. Collect candidate R2 keys from the rows being deleted.
    try:
        rows = svc.table("submissions").select("id,data").in_("id", ids).execute().data or []
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch submissions: {exc}",
        ) from exc

    candidate_keys: set[str] = set()
    for row in rows:
        candidate_keys |= _submission_keys(row.get("data") or {})

    # 2. Delete the DB rows.
    try:
        svc.table("submissions").delete().in_("id", ids).execute()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not delete submissions: {exc}",
        ) from exc

    # 3. Best-effort R2 cleanup — only keys no surviving submission references.
    if settings.r2_enabled and candidate_keys:
        try:
            survivors = svc.table("submissions").select("data").execute().data or []
            still_referenced: set[str] = set()
            for s in survivors:
                still_referenced |= _submission_keys(s.get("data") or {})
            delete_keys(list(candidate_keys - still_referenced))
        except Exception:
            pass  # best-effort — never fail the delete on R2 cleanup

    return {"deleted": len(rows)}


# ---------------------------------------------------------------------------
# New-artist queue
# ---------------------------------------------------------------------------


class NewArtistUpdateBody(BaseModel):
    spotify_url: str = ""
    apple_music_url: str = ""


class AdminCreateUser(BaseModel):
    email: str
    password: str
    full_name: str = ""
    artist_name: str = ""
    phone: str = ""
    plan: str = "free"


class AdminSetPassword(BaseModel):
    password: str


class AdminUserUpdate(BaseModel):
    full_name: str | None = None
    artist_name: str | None = None
    phone: str | None = None
    city: str | None = None
    state: str | None = None
    date_of_birth: str | None = None
    gender: str | None = None
    bio: str | None = None
    spotify_url: str | None = None
    apple_music_url: str | None = None
    instagram: str | None = None
    youtube_url: str | None = None
    custom_label_name: str | None = None
    plan: str | None = None
    expires_at: str | None = None  # "YYYY-MM-DD" to set; "" to clear to NULL; None = no change


@router.get("/new-artist-queue", dependencies=[Depends(_require_admin)])
async def list_new_artist_queue() -> dict:
    """Return all new-artist queue entries — pending first."""
    svc = get_service_client()
    try:
        resp = (
            svc.table("new_artist_queue")
            .select("*")
            .order("status", desc=False)   # pending before updated
            .order("created_at", desc=True)
            .execute()
        )
    except Exception as exc:
        # Table may not exist yet — return empty list gracefully instead of 502
        if "PGRST205" in str(exc) or "schema cache" in str(exc).lower() or "does not exist" in str(exc).lower():
            return {"entries": [], "hint": "Run migration 0004_apple_music_and_new_artist_queue.sql in Supabase SQL Editor to enable this feature."}
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"Could not fetch queue: {exc}")
    return {"entries": resp.data or []}


@router.patch("/new-artist-queue/{entry_id}", dependencies=[Depends(_require_admin)])
async def update_new_artist(entry_id: str, body: NewArtistUpdateBody) -> dict:
    """Save Spotify + Apple Music links for a queued new artist and update their profile."""
    svc = get_service_client()

    try:
        entry_resp = svc.table("new_artist_queue").select("*").eq("id", entry_id).limit(1).execute()
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"DB error: {exc}")
    if not entry_resp.data:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Entry not found")

    entry = entry_resp.data[0]
    user_email = entry.get("user_email", "")

    # Look up user UUID by email so we can update their profile.
    user_id: str | None = None
    try:
        all_users = svc.auth.admin.list_users()
        match = next((u for u in all_users if getattr(u, "email", "") == user_email), None)
        if match:
            user_id = str(match.id)
    except Exception:
        pass

    if user_id:
        try:
            profile_service.upsert_profile(user_id, {
                "spotify_url": body.spotify_url,
                "apple_music_url": body.apple_music_url,
            })
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"Profile update failed: {exc}")

    try:
        upd = (
            svc.table("new_artist_queue")
            .update({
                "spotify_url": body.spotify_url,
                "apple_music_url": body.apple_music_url,
                "status": "updated",
                "updated_at": datetime.now(timezone.utc).isoformat(),
            })
            .eq("id", entry_id)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"Queue update failed: {exc}")

    return upd.data[0] if upd.data else {}


# ---------------------------------------------------------------------------
# User create / edit / delete
# ---------------------------------------------------------------------------


def _is_duplicate_email(msg: str) -> bool:
    m = msg.lower()
    return any(x in m for x in ("already registered", "already exists", "duplicate", "unique constraint", "already been registered"))


@router.post("/users", dependencies=[Depends(_require_admin)], status_code=201)
async def admin_create_user(body: AdminCreateUser) -> dict:
    """Create a pre-confirmed user, bypassing email verification."""
    if not body.email or not body.password:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Email and password are required")
    if len(body.password) < 6:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Password must be at least 6 characters")

    svc = get_service_client()
    try:
        result = await run_in_threadpool(
            svc.auth.admin.create_user,
            {
                "email": body.email,
                "password": body.password,
                "email_confirm": True,
                "user_metadata": {
                    k: v for k, v in {
                        "full_name": body.full_name,
                        "artist_name": body.artist_name,
                        "phone": body.phone,
                    }.items() if v
                },
            },
        )
    except Exception as exc:
        msg = str(exc)
        if _is_duplicate_email(msg):
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="An account with this email already exists")
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=msg) from exc

    user = result.user
    uid = str(user.id)

    # Upsert profile fields (handle_new_user trigger already created the blank row)
    profile_fields = {k: v for k, v in {
        "full_name": body.full_name,
        "artist_name": body.artist_name,
        "phone": body.phone,
    }.items() if v}
    if profile_fields:
        try:
            profile_service.upsert_profile(uid, profile_fields)
        except Exception as exc:
            _log.warning("Could not upsert profile for admin-created user %s: %s", uid, exc)

    # Assign non-free plan if requested (handle_new_user trigger already made the Free row)
    if body.plan and body.plan != "free":
        try:
            plan_obj = Plan(body.plan)
            assign_plan(uid, plan_obj)
            # This is the user's first-ever plan grant — inherently a new activation.
            credit_referral(referred_user_id=uid, plan=plan_obj, source="admin")
        except Exception as exc:
            _log.warning("Could not assign plan for admin-created user %s: %s", uid, exc)

    return {"created": True, "user_id": uid, "email": body.email}


@router.patch("/users/{user_id}", dependencies=[Depends(_require_admin)])
async def update_user(user_id: str, body: AdminUserUpdate) -> dict:
    svc = get_service_client()
    plan_changed: Plan | None = None
    new_expiry: str | None = None
    expires_at_updated = False
    try:
        # 1. Validate + apply plan change first (before touching profile/meta)
        if body.plan is not None:
            try:
                plan_changed = Plan(body.plan)
            except ValueError:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid plan '{body.plan}'. Valid values: {[p.value for p in Plan]}",
                )
            # Snapshot the plan BEFORE overwriting it, so we only credit a referral
            # commission on a genuine change — an admin resaving unrelated profile
            # fields still sends the same `plan` value and must not double-credit.
            prev = (
                svc.table("subscriptions").select("plan").eq("user_id", user_id).limit(1).execute()
            )
            prev_rows = prev.data or []
            prev_plan = prev_rows[0].get("plan") if prev_rows else None

            assign_plan(user_id, plan_changed)
            if prev_plan != plan_changed.value:
                credit_referral(referred_user_id=user_id, plan=plan_changed, source="admin")

        # 2. Apply explicit expires_at override — runs AFTER assign_plan so it wins
        #    when both fields arrive in the same request (assign_plan always sets
        #    now+365 for paid plans; this overwrites that default).
        #    None   = field absent → skip entirely (no subscriptions write)
        #    ""     = clear to NULL (plan never expires — complimentary)
        #    "YYYY-MM-DD" = store as end-of-day UTC
        if body.expires_at is not None:
            new_expiry = _parse_expires_at(body.expires_at)
            (
                svc.table("subscriptions")
                .update({"expires_at": new_expiry})
                .eq("user_id", user_id)
                .execute()
            )
            expires_at_updated = True

        # 3. Upsert ALL profile fields into the profiles table.
        # This includes full_name/artist_name/phone — profiles is the authoritative
        # store; auth.user_metadata is synced separately as a best-effort step.
        # Skip None and "" — Postgres date columns reject empty strings,
        # and we don't want to clear fields the admin left blank.
        # Exclude plan and expires_at — both belong to subscriptions, not profiles.
        profile_fields = {
            k: v for k, v in body.model_dump().items()
            if k not in ("plan", "expires_at")
            and v is not None
            and v != ""
        }
        if profile_fields:
            profile_service.upsert_profile(user_id, profile_fields)

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not update user: {exc}",
        ) from exc

    # 3. Best-effort sync of name fields to auth.user_metadata so JWT claims
    # stay current. Failure here is non-fatal — profile table is already saved.
    meta = {k: v for k, v in {
        "full_name": body.full_name,
        "artist_name": body.artist_name,
        "phone": body.phone,
    }.items() if v is not None and v != ""}
    if meta:
        try:
            svc.auth.admin.update_user_by_id(user_id, {"user_metadata": meta})
        except Exception as exc:
            _log.warning("Could not sync user_metadata for %s: %s", user_id, exc)

    result: dict = {"updated": True, "user_id": user_id}
    if plan_changed is not None:
        spec = get_spec(plan_changed)
        result["plan"] = plan_changed.value
        result["plan_name"] = spec.name
    if expires_at_updated:
        result["expires_at"] = new_expiry  # None when cleared, ISO string when set
    return result


@router.delete("/users/{user_id}", dependencies=[Depends(_require_admin)])
async def delete_user_endpoint(user_id: str) -> dict:
    svc = get_service_client()
    try:
        svc.auth.admin.delete_user(user_id)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not delete user: {exc}",
        ) from exc
    return {"deleted": True, "user_id": user_id}


@router.patch("/users/{user_id}/password", dependencies=[Depends(_require_admin)])
async def set_user_password(user_id: str, body: AdminSetPassword) -> dict:
    """Override a user's password — admin only. Plaintext passwords are never
    stored; this sets a new bcrypt hash via Supabase GoTrue admin API."""
    if len(body.password) < 6:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must be at least 6 characters",
        )
    svc = get_service_client()
    try:
        await run_in_threadpool(
            svc.auth.admin.update_user_by_id,
            user_id,
            {"password": body.password},
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not update password: {exc}",
        ) from exc
    return {"updated": True}


# ---------------------------------------------------------------------------
# Plan purchases
# ---------------------------------------------------------------------------


@router.get("/purchases", dependencies=[Depends(_require_admin)])
async def list_purchases() -> dict:
    """Return all confirmed (paid) plan purchases with user details.

    Fetches every subscription where plan != 'free', joins with auth user data,
    and computes total active revenue for the stats panel.
    """
    svc = get_service_client()
    try:
        subs_resp = (
            svc.table("subscriptions")
            .select("*")
            .neq("plan", "free")
            .order("started_at", desc=True)
            .execute()
        )
        subs = subs_resp.data or []
        raw_users = _fetch_all_users(svc)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch purchases: {exc}",
        ) from exc

    user_map: dict[str, dict] = {}
    for u in raw_users:
        uid = str(getattr(u, "id", "") or "")
        user_meta: dict = getattr(u, "user_metadata", None) or {}
        user_map[uid] = {
            "email": getattr(u, "email", "") or "",
            "full_name": user_meta.get("full_name", "") or "",
            "artist_name": user_meta.get("artist_name", "") or "",
        }

    purchases = []
    plan_counts: dict[str, int] = {}
    total_revenue = 0

    for sub in subs:
        uid = sub.get("user_id", "")
        user = user_map.get(uid, {})
        plan_key = sub.get("plan") or ""
        if not plan_key or plan_key == "free":
            continue
        plan_name = _PLAN_NAMES.get(plan_key, plan_key.replace("-", " ").replace("_", " ").title())
        plan_price = _PLAN_PRICES_INR.get(plan_key, 0)

        if sub.get("status") == "active":
            plan_counts[plan_key] = plan_counts.get(plan_key, 0) + 1
            total_revenue += plan_price

        purchases.append({
            "id": sub.get("id"),
            "user_id": uid,
            "email": user.get("email", ""),
            "full_name": user.get("full_name", ""),
            "artist_name": user.get("artist_name", ""),
            "plan": plan_key,
            "plan_name": plan_name,
            "plan_price_inr": plan_price,
            "status": sub.get("status", ""),
            "payment_ref": sub.get("payment_ref"),
            "started_at": sub.get("started_at"),
            "expires_at": sub.get("expires_at"),
        })

    return {
        "purchases": purchases,
        "total": len(purchases),
        "plan_counts": plan_counts,
        "total_revenue_inr": total_revenue,
    }


# ---------------------------------------------------------------------------
# Home content management
# ---------------------------------------------------------------------------

_ALLOWED_IMAGE_TYPES: dict[str, str] = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
}
_MAX_IMAGE_BYTES = 5 * 1024 * 1024  # 5 MB


@router.get("/home", dependencies=[Depends(_require_admin)])
async def admin_get_home() -> HomeContent:
    try:
        return home_service.get_home_content()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch home content: {exc}",
        ) from exc


@router.put("/home", dependencies=[Depends(_require_admin)])
async def admin_update_home(body: HomeContent) -> HomeContent:
    try:
        return home_service.upsert_home_content(body)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not save home content: {exc}",
        ) from exc


@router.post("/home/artist-image", dependencies=[Depends(_require_admin)])
async def upload_artist_image(file: UploadFile = File(...)) -> dict[str, str]:
    if not settings.r2_enabled:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="R2 storage is not configured. Set R2_ACCOUNT_ID, R2_ACCESS_KEY_ID, R2_SECRET_ACCESS_KEY, and R2_BUCKET_NAME.",
        )
    if file.content_type not in _ALLOWED_IMAGE_TYPES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only JPEG, PNG, or WebP images are allowed.",
        )
    data = await file.read()
    if len(data) > _MAX_IMAGE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File size must not exceed 5 MB.",
        )
    ext = _ALLOWED_IMAGE_TYPES[file.content_type]
    key = f"home/artists/{uuid4().hex}{ext}"
    content_type = file.content_type
    try:
        await run_in_threadpool(upload_bytes, key, data, content_type)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Image upload to R2 failed: {exc}",
        ) from exc
    base = settings.oauth_callback_base_url.rstrip("/")
    return {"url": f"{base}/home/assets/{key}"}


# ---------------------------------------------------------------------------
# Media download (R2 presigned GET)
# ---------------------------------------------------------------------------


@router.get("/media/download-url", dependencies=[Depends(_require_admin)])
async def get_download_url(key: str = Query(...)) -> dict:
    """Generate a 15-minute presigned GET URL so the admin can download an R2 file."""
    if not settings.r2_enabled:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="R2 storage not configured.",
        )
    if not key.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="key is required.")
    try:
        url = presign_get(key.strip(), expires_in=900)
        return {"url": url, "expires_in": 900}
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not generate download URL: {exc}",
        ) from exc


# ---------------------------------------------------------------------------
# Earnings / song_stats admin CRUD
# ---------------------------------------------------------------------------

_VALID_MONTHS = {
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
}

_PLATFORM_GROUP_MAP: dict[str, str] = {
    "spotify": "Spotify",
    "apple music": "Apple Music",
    "youtube": "YouTube",
    "youtube music": "YouTube",
    "facebook": "Facebook",
    "meta": "Facebook",
    "facebook/meta": "Facebook",
    "amazon": "Amazon",
    "amazon music": "Amazon",
    "jiosaavn": "JioSaavn",
    "gaana": "Gaana",
    "tiktok": "TikTok",
}


def _derive_platform_group(platform: str) -> str:
    return _PLATFORM_GROUP_MAP.get(platform.lower().strip(), "Other")


class AdminPlatformEntry(BaseModel):
    platform: str
    streams: int = Field(ge=0)
    revenue: str  # Decimal string — validated server-side


class AdminSongStatBulkCreate(BaseModel):
    user_email: str
    song_title: str
    artist_name: str
    period_month: str
    period_year: int
    submission_id: Optional[str] = None
    entries: List[AdminPlatformEntry]


class AdminSongStatUpdate(BaseModel):
    streams: Optional[int] = Field(None, ge=0)
    revenue: Optional[str] = None  # Decimal string ≥ 0


class AdminBalanceAdjustmentCreate(BaseModel):
    user_email: str
    amount: str  # signed Decimal string — positive = credit, negative = debit
    comment: str = Field(min_length=1, max_length=500)


def _parse_revenue(rev: str) -> Decimal:
    try:
        d = Decimal(str(rev))
        if d < 0:
            raise ValueError
        return d
    except (InvalidOperation, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"revenue must be a non-negative decimal string, got: {rev!r}",
        ) from exc


@router.get("/artist-earnings", dependencies=[Depends(_require_admin)])
async def get_artist_earnings(
    email: str = Query(default=""),
    q: str = Query(default=""),
) -> dict:
    """Return all song_stats rows + balance for one artist.

    Pass ?email= for an exact lookup.
    Pass ?q= to search by artist_name in profiles (returns a match list for the
    admin to select from, then re-query with ?email=).
    """
    svc = get_service_client()

    # Name search mode — return candidate list only.
    if not email and q:
        try:
            res = (
                svc.table("profiles")
                .select("user_id,artist_name")
                .ilike("artist_name", f"%{q}%")
                .limit(20)
                .execute()
            )
            candidates = []
            for p in (res.data or []):
                # Resolve email from auth users via user_id
                try:
                    u = svc.auth.admin.get_user_by_id(p["user_id"])
                    candidates.append({
                        "email": getattr(u.user, "email", "") or "",
                        "artist_name": p.get("artist_name") or "",
                    })
                except Exception:
                    pass
            return {"candidates": candidates}
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                                detail=str(exc)) from exc

    if not email:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Provide ?email= or ?q=")

    email = email.lower().strip()

    # Fetch all song_stats rows including id (admin view).
    rows: list[dict] = []
    start = 0
    try:
        while True:
            resp = (
                svc.table("song_stats")
                .select("id,user_email,song_title,artist_name,platform,platform_group,"
                        "period_month,period_year,streams,revenue,submission_id")
                .eq("user_email", email)
                .range(start, start + 999)
                .execute()
            )
            batch = resp.data or []
            rows.extend(batch)
            if len(batch) < 1000:
                break
            start += 1000
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    # Balance row.
    artist_info: dict[str, Any] = {"email": email, "artist_name": ""}
    try:
        bal_res = (
            svc.table("artist_balances")
            .select("total_earned,total_withdrawn,available_balance,last_updated")
            .eq("user_email", email)
            .maybe_single()
            .execute()
        )
        b = bal_res.data or {}
        artist_info.update({
            "total_earned": float(Decimal(str(b.get("total_earned") or 0)).quantize(Decimal("0.01"))),
            "total_withdrawn": float(Decimal(str(b.get("total_withdrawn") or 0)).quantize(Decimal("0.01"))),
            "available_balance": float(Decimal(str(b.get("available_balance") or 0)).quantize(Decimal("0.01"))),
            "last_updated": b.get("last_updated"),
        })
    except Exception:
        artist_info.update({"total_earned": 0, "total_withdrawn": 0,
                            "available_balance": 0, "last_updated": None})

    # Artist display name — use first song_stats row (already loaded).
    if rows:
        artist_info["artist_name"] = rows[0].get("artist_name") or ""

    # Normalise revenue to float for display.
    for r in rows:
        r["revenue"] = float(Decimal(str(r.get("revenue") or 0)).quantize(Decimal("0.01")))

    # Manual balance adjustments (migration 0012) for this artist's Adjust
    # Balance panel — degrades to [] if the table doesn't exist yet.
    adjustments: list[dict] = []
    try:
        adj_res = (
            svc.table("balance_adjustments")
            .select("id,amount,comment,created_at")
            .eq("user_email", email)
            .order("created_at", desc=True)
            .execute()
        )
        for a in (adj_res.data or []):
            a["amount"] = float(Decimal(str(a.get("amount") or 0)).quantize(Decimal("0.01")))
            adjustments.append(a)
    except Exception:
        pass

    return {"artist": artist_info, "rows": rows, "adjustments": adjustments}


@router.post("/song-stats", dependencies=[Depends(_require_admin)])
async def admin_add_song_stats(body: AdminSongStatBulkCreate) -> dict:
    """Bulk-add platform rows for one song × month.

    Uses upsert on the UNIQUE key so re-posting the same data is idempotent
    (updates streams/revenue instead of raising a duplicate error).
    """
    if body.period_month not in _VALID_MONTHS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"period_month must be a full English month name, got: {body.period_month!r}",
        )
    current_year = datetime.now(timezone.utc).year
    if not (1900 <= body.period_year <= current_year + 1):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"period_year out of range: {body.period_year}",
        )
    if not body.entries:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="entries must contain at least one platform row",
        )

    # Validate all revenues before any write.
    validated_revenues = [_parse_revenue(e.revenue) for e in body.entries]

    email = body.user_email.lower().strip()
    now = datetime.now(timezone.utc).isoformat()

    rows_to_insert = []
    for entry, rev in zip(body.entries, validated_revenues):
        rows_to_insert.append({
            "user_email": email,
            "song_title": body.song_title,
            "artist_name": body.artist_name,
            "platform": entry.platform,
            "platform_group": _derive_platform_group(entry.platform),
            "period_month": body.period_month,
            "period_year": body.period_year,
            "streams": entry.streams,
            "revenue": str(rev),
            "submission_id": body.submission_id or None,
            "updated_at": now,
        })

    svc = get_service_client()
    try:
        res = svc.table("song_stats").upsert(
            rows_to_insert,
            on_conflict="user_email,song_title,platform,period_month,period_year",
        ).execute()
        inserted_rows = res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    balance = recompute_balance(email)
    return {"inserted": len(inserted_rows), "rows": inserted_rows, "balance": balance}


@router.patch("/song-stats/{row_id}", dependencies=[Depends(_require_admin)])
async def admin_update_song_stat(row_id: str, body: AdminSongStatUpdate) -> dict:
    """Update streams and/or revenue on an existing song_stats row.

    Key fields (song_title, platform, period_month, period_year) are immutable
    via PATCH — delete and re-add to correct them.
    """
    svc = get_service_client()

    # Fetch the row first to get user_email for balance recompute.
    # Plain select + limit(1) instead of .maybe_single() — the latter raises
    # (PostgREST 406) rather than returning an empty result when the id
    # doesn't match, which would surface as a confusing 502 here instead of
    # the intended 404 below.
    try:
        existing_res = (
            svc.table("song_stats")
            .select("id,user_email,song_title,artist_name,platform,platform_group,"
                    "period_month,period_year,streams,revenue,submission_id")
            .eq("id", row_id)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    existing_rows = existing_res.data or []
    if not existing_rows:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND,
                            detail=f"song_stats row {row_id!r} not found")

    existing_row = existing_rows[0]
    email = existing_row["user_email"]
    patch: dict[str, Any] = {"updated_at": datetime.now(timezone.utc).isoformat()}

    if body.streams is not None:
        patch["streams"] = body.streams
    if body.revenue is not None:
        patch["revenue"] = str(_parse_revenue(body.revenue))

    if len(patch) == 1:  # only updated_at — nothing to change
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                            detail="Provide at least one of: streams, revenue")

    try:
        res = svc.table("song_stats").update(patch).eq("id", row_id).execute()
        updated_row = (res.data or [existing.data])[0]
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    updated_row["revenue"] = float(
        Decimal(str(updated_row.get("revenue") or 0)).quantize(Decimal("0.01"))
    )
    balance = recompute_balance(email)
    return {"row": updated_row, "balance": balance}


@router.delete("/song-stats/{row_id}", dependencies=[Depends(_require_admin)])
async def admin_delete_song_stat(row_id: str) -> dict:
    """Delete a song_stats row and recompute the artist's balance."""
    svc = get_service_client()

    try:
        existing_res = (
            svc.table("song_stats")
            .select("user_email")
            .eq("id", row_id)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    existing_rows = existing_res.data or []
    if not existing_rows:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND,
                            detail=f"song_stats row {row_id!r} not found")

    email = existing_rows[0]["user_email"]

    try:
        svc.table("song_stats").delete().eq("id", row_id).execute()
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    balance = recompute_balance(email)
    return {"deleted": True, "balance": balance}


@router.post("/balance-adjustments", dependencies=[Depends(_require_admin)])
async def admin_create_balance_adjustment(body: AdminBalanceAdjustmentCreate) -> dict:
    """Manually credit or debit an artist's balance with a required comment.

    Writes an immutable row to balance_adjustments (migration 0012) rather
    than touching artist_balances directly — recompute_balance() sums this
    ledger every time it runs, so the adjustment survives every future
    recompute (song_stats edits, monthly royalty ingest) instead of being
    silently overwritten.
    """
    try:
        amount = Decimal(str(body.amount))
    except (InvalidOperation, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"amount must be a decimal string, got: {body.amount!r}",
        ) from exc
    if amount == 0:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                            detail="amount must be nonzero")

    email = body.user_email.lower().strip()

    # recompute_balance() floors available_balance at 0 but NOT total_earned —
    # a debit larger than the artist's current total_earned would otherwise
    # drive it negative. Reject before writing anything.
    current = get_balance(email)
    if amount < 0 and abs(amount) > Decimal(str(current["total_earned"])):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(f"Debit of {abs(amount)} exceeds current total earned "
                    f"of {current['total_earned']}"),
        )

    svc = get_service_client()
    row = {
        "user_email": email,
        "amount": str(amount),
        "comment": body.comment,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    try:
        res = svc.table("balance_adjustments").insert(row).execute()
        created = (res.data or [row])[0]
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    balance = recompute_balance(email)
    return {"ok": True, "balance": balance, "adjustment": created}


@router.delete("/balance-adjustments/{adjustment_id}", dependencies=[Depends(_require_admin)])
async def admin_delete_balance_adjustment(adjustment_id: str) -> dict:
    """Delete a manual adjustment and recompute the artist's balance.

    No separate credit-back step is needed (unlike withdrawal_requests) —
    recompute_balance() re-sums balance_adjustments from scratch, so removing
    a row automatically reverses its effect on total_earned/available_balance.
    """
    svc = get_service_client()

    # Plain select + limit(1), not .maybe_single() — PostgREST 406s on zero
    # rows under the singular Accept header, which would surface as a
    # confusing 502 here instead of the intended 404 below.
    try:
        existing_res = (
            svc.table("balance_adjustments")
            .select("user_email")
            .eq("id", adjustment_id)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    existing_rows = existing_res.data or []
    if not existing_rows:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND,
                            detail=f"balance_adjustments row {adjustment_id!r} not found")

    email = existing_rows[0]["user_email"]

    try:
        svc.table("balance_adjustments").delete().eq("id", adjustment_id).execute()
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    balance = recompute_balance(email)
    return {"deleted": True, "balance": balance}


@router.get("/song-stats/submissions/{email}", dependencies=[Depends(_require_admin)])
async def admin_list_artist_submissions(email: str) -> dict:
    """Return a lightweight list of an artist's submissions for the Add modal dropdown."""
    svc = get_service_client()
    email = email.lower().strip()
    try:
        res = (
            svc.table("submissions")
            .select("id,type,status,data")
            .eq("user_email", email)
            .order("created_at", desc=True)
            .execute()
        )
        rows = res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY,
                            detail=str(exc)) from exc

    submissions = []
    for r in rows:
        data = r.get("data") or {}
        title = (data.get("song_title") or data.get("album_name") or
                 data.get("song_name") or "")
        submissions.append({
            "id": r["id"],
            "title": title,
            "type": r.get("type") or "",
            "status": r.get("status") or "",
        })
    return {"submissions": submissions}


# ---------------------------------------------------------------------------
# Central USD -> INR rate + bulk Excel import into song_stats
# ---------------------------------------------------------------------------

_MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
_MAX_IMPORT_ROWS = 20_000
_REQUIRED_IMPORT_HEADERS = ("song", "streams", "revenue", "month", "year", "platform")


def _norm_title(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _to_int(v: Any) -> int:
    return int(float(v))


class ImportRowError(ValueError):
    """Raised for a malformed import header/row — message is 422-ready."""


def _parse_import_rows(header_row: tuple, data_rows: Any) -> dict[str, Any]:
    """Pure parsing + aggregation of an Excel report's rows.

    No Supabase, no file I/O, no FastAPI — takes exactly what
    ``openpyxl``'s ``iter_rows(values_only=True)`` yields (the header tuple,
    then an iterable of data-row tuples) and returns the per
    (song, platform, month, year) aggregated groups, summed across any
    ``Country`` column since song_stats has no country grain. Raises
    ``ImportRowError`` on any bad header/row — kept separate from
    HTTPException so this function has zero web-framework dependency and can
    be unit-tested with plain tuples.
    """
    if not header_row:
        raise ImportRowError("Workbook has no header row")

    col_index: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        col_index[str(cell).strip().lower()] = i

    missing = [h for h in _REQUIRED_IMPORT_HEADERS if h not in col_index]
    if missing:
        raise ImportRowError(f"Missing required column(s): {', '.join(missing)}")

    has_artist_col = "artistname" in col_index

    def _cell(row: tuple, key: str) -> Any:
        idx = col_index.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    groups: dict[tuple[str, str, str, int], dict[str, Any]] = {}
    file_artist_names_norm: set[str] = set()
    first_file_artist_name = ""
    row_count = 0
    for row in data_rows:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        song_title = str(_cell(row, "song") or "").strip()
        month = str(_cell(row, "month") or "").strip().title()
        platform_raw = str(_cell(row, "platform") or "").strip()

        if not song_title and not month and not platform_raw:
            # Trailing subtotal/grand-total row that some DSP report exports
            # append (only Streams/Revenue filled in, every identifying
            # column blank). Not real per-song data — skip like a blank row
            # rather than hard-failing the whole import on it.
            continue

        row_count += 1
        if row_count > _MAX_IMPORT_ROWS:
            raise ImportRowError(f"Workbook exceeds the {_MAX_IMPORT_ROWS}-row limit")

        excel_row_num = row_count + 1  # +1 for the header row

        if not song_title:
            raise ImportRowError(f"Row {excel_row_num}: missing Song")
        if month not in _VALID_MONTHS:
            raise ImportRowError(f"Row {excel_row_num}: invalid Month {month!r}")
        try:
            streams = _to_int(_cell(row, "streams"))
            year = _to_int(_cell(row, "year"))
            revenue_usd = Decimal(str(_cell(row, "revenue")))
        except (TypeError, ValueError, InvalidOperation):
            raise ImportRowError(f"Row {excel_row_num}: Streams/Revenue/Year must be numeric")
        if streams < 0 or revenue_usd < 0:
            raise ImportRowError(f"Row {excel_row_num}: Streams/Revenue must be non-negative")

        if has_artist_col:
            a_raw = str(_cell(row, "artistname") or "").strip()
            if a_raw:
                file_artist_names_norm.add(_norm_title(a_raw))
                if not first_file_artist_name:
                    first_file_artist_name = a_raw

        # Group by the CANONICAL platform, not the raw label — DSP reports
        # often split one platform into multiple raw sub-lines (e.g. "YouTube"
        # and "YouTube (PDL)" both normalize to "YouTube"). Grouping on the
        # raw string would leave those as separate groups that later collide
        # on the same (song_title, platform, month, year) upsert key, which
        # Postgres rejects with "ON CONFLICT DO UPDATE command cannot affect
        # row a second time" — merge them here instead so the upsert batch
        # never has two rows for the same conflict key.
        platform, platform_group = normalize_platform(platform_raw)
        key = (_norm_title(song_title), platform, month, year)
        g = groups.get(key)
        if g is None:
            g = {
                "song_title": song_title, "platform_raw": platform_raw,
                "platform": platform, "platform_group": platform_group,
                "month": month, "year": year,
                "streams": 0, "revenue_usd": Decimal("0"),
            }
            groups[key] = g
        g["streams"] += streams
        g["revenue_usd"] += revenue_usd

    if not groups:
        raise ImportRowError("No data rows found in workbook")

    return {
        "groups": groups,
        "file_artist_names_norm": file_artist_names_norm,
        "first_file_artist_name": first_file_artist_name,
    }


class FxRateUpdate(BaseModel):
    usd_to_inr: str  # Decimal string — validated server-side


@router.get("/fx-rate", dependencies=[Depends(_require_admin)])
async def get_fx_rate() -> dict:
    """Return the centrally-set USD -> INR rate used by /admin/song-stats/import.

    Not per-artist, not per-upload — set once here, then every subsequent
    import for every artist uses whatever is currently stored.
    """
    svc = get_service_client()
    try:
        res = svc.table("fx_rate_settings").select("usd_to_inr,updated_at").eq("id", 1).limit(1).execute()
        rows = res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc
    if not rows:
        return {"usd_to_inr": None, "updated_at": None}
    row = rows[0]
    return {
        "usd_to_inr": float(Decimal(str(row["usd_to_inr"]))),
        "updated_at": row.get("updated_at"),
    }


@router.put("/fx-rate", dependencies=[Depends(_require_admin)])
async def update_fx_rate(body: FxRateUpdate) -> dict:
    """Set/replace the centrally-stored USD -> INR conversion rate."""
    rate = _parse_revenue(body.usd_to_inr)
    if rate <= 0:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="usd_to_inr must be greater than 0",
        )
    svc = get_service_client()
    now = datetime.now(timezone.utc).isoformat()
    try:
        res = svc.table("fx_rate_settings").upsert(
            {"id": 1, "usd_to_inr": str(rate), "updated_at": now},
            on_conflict="id",
        ).execute()
        row = (res.data or [{}])[0]
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc
    return {
        "usd_to_inr": float(Decimal(str(row.get("usd_to_inr") or rate))),
        "updated_at": row.get("updated_at") or now,
    }


@router.post("/song-stats/import", dependencies=[Depends(_require_admin)])
async def admin_import_song_stats(
    file: UploadFile = File(...),
    email: str = Form(...),
) -> dict:
    """Bulk-import one artist's monthly Excel royalty report into song_stats.

    Pure deterministic column parsing + matching — no AI/LLM involved anywhere.
    The USD -> INR rate always comes from the centrally-stored fx_rate_settings
    row (GET/PUT /admin/fx-rate) — it is never re-entered per upload.
    """
    svc = get_service_client()
    email = email.lower().strip()

    # Fail fast, before touching the file or writing anything, if no rate has
    # ever been set — this is what makes "just upload, you're sorted" safe.
    try:
        rate_res = svc.table("fx_rate_settings").select("usd_to_inr").eq("id", 1).limit(1).execute()
        rate_rows = rate_res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc
    if not rate_rows:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"error": "usd_to_inr_not_set", "detail": "Set the USD→INR rate first."},
        )
    usd_to_inr = Decimal(str(rate_rows[0]["usd_to_inr"]))

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="File must be a .xlsx workbook")

    raw_bytes = await file.read()
    if len(raw_bytes) > _MAX_IMPORT_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File exceeds the {_MAX_IMPORT_FILE_BYTES // (1024 * 1024)} MB limit",
        )

    try:
        wb = openpyxl.load_workbook(BytesIO(raw_bytes), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail=f"Could not read workbook: {exc}") from exc

    try:
        parsed = _parse_import_rows(header_row, rows_iter)
    except ImportRowError as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                            detail=str(exc)) from exc

    groups = parsed["groups"]
    file_artist_names_norm = parsed["file_artist_names_norm"]
    first_file_artist_name = parsed["first_file_artist_name"]

    # This artist's existing song_stats, fetched once (not per-row).
    try:
        existing_res = (
            svc.table("song_stats")
            .select("song_title,artist_name,submission_id")
            .eq("user_email", email)
            .execute()
        )
        existing_rows = existing_res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc

    existing_titles: set[str] = set()
    title_to_submission: dict[str, Optional[str]] = {}
    artist_name_on_file = ""
    for r in existing_rows:
        norm = _norm_title(r.get("song_title") or "")
        existing_titles.add(norm)
        if norm not in title_to_submission:
            title_to_submission[norm] = r.get("submission_id")
        if not artist_name_on_file and r.get("artist_name"):
            artist_name_on_file = r["artist_name"]

    # Best-effort submission_id backfill for brand-new songs — same title
    # fallback admin_list_artist_submissions already uses.
    needs_submission_lookup = any(
        _norm_title(g["song_title"]) not in existing_titles for g in groups.values()
    )
    submission_title_map: dict[str, str] = {}
    if needs_submission_lookup:
        try:
            sub_res = (
                svc.table("submissions")
                .select("id,data")
                .eq("user_email", email)
                .execute()
            )
            for r in (sub_res.data or []):
                d = r.get("data") or {}
                title = d.get("song_title") or d.get("album_name") or d.get("song_name") or ""
                if title:
                    submission_title_map.setdefault(_norm_title(title), r["id"])
        except Exception:
            submission_title_map = {}

    artist_name_for_insert = artist_name_on_file or first_file_artist_name

    now = datetime.now(timezone.utc).isoformat()
    rows_to_upsert: list[dict[str, Any]] = []
    songs_new: set[str] = set()
    songs_updated: set[str] = set()
    total_streams = 0
    total_revenue_inr = Decimal("0")

    for g in groups.values():
        norm_title = _norm_title(g["song_title"])
        is_new = norm_title not in existing_titles
        (songs_new if is_new else songs_updated).add(norm_title)

        submission_id = title_to_submission.get(norm_title) or submission_title_map.get(norm_title)
        platform, platform_group = g["platform"], g["platform_group"]
        revenue_inr = g["revenue_usd"] * usd_to_inr

        total_streams += g["streams"]
        total_revenue_inr += revenue_inr

        rows_to_upsert.append({
            "user_email": email,
            "song_title": g["song_title"],
            "artist_name": artist_name_for_insert,
            "platform": platform,
            "platform_group": platform_group,
            "period_month": g["month"],
            "period_year": g["year"],
            "streams": g["streams"],
            "revenue": str(revenue_inr),
            "submission_id": submission_id,
            "updated_at": now,
        })

    try:
        res = svc.table("song_stats").upsert(
            rows_to_upsert,
            on_conflict="user_email,song_title,platform,period_month,period_year",
        ).execute()
        upserted_rows = res.data or []
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=str(exc)) from exc

    balance = recompute_balance(email)

    warnings: list[str] = []
    if artist_name_on_file and file_artist_names_norm:
        if _norm_title(artist_name_on_file) not in file_artist_names_norm:
            warnings.append(
                f"File's ArtistName does not match the artist on record for {email} "
                f"({artist_name_on_file!r}) — double-check this is the right artist."
            )

    return {
        "imported": {
            "songs_new": len(songs_new),
            "songs_updated": len(songs_updated),
            "rows_written": len(rows_to_upsert),
            "total_streams": total_streams,
            "total_revenue_inr": float(total_revenue_inr.quantize(Decimal("0.01"))),
        },
        "rows": upserted_rows,
        "balance": balance,
        "warnings": warnings,
    }


# ---------------------------------------------------------------------------
# Tunefry Daily (blog)
# ---------------------------------------------------------------------------

_BLOG_REVIEW_ERROR_DETAIL: dict[str, str] = {
    "not_found": "Post not found.",
    "final_copy_required": "Both a final title and final body are required to approve.",
    "admin_note_required": "A comment is required to decline.",
    "not_approved": "Only approved posts can be featured/popular.",
}


def _blog_review_error_to_http(exc: ReviewError) -> HTTPException:
    code = status.HTTP_404_NOT_FOUND if str(exc) == "not_found" else status.HTTP_400_BAD_REQUEST
    detail = _BLOG_REVIEW_ERROR_DETAIL.get(str(exc), str(exc))
    return HTTPException(status_code=code, detail=detail)


def _notify_artist_blog_review(email: str, title: str, body: str) -> None:
    if not email:
        return
    try:
        get_service_client().table("notifications").insert(
            {"title": title, "body": body, "recipient_email": email}
        ).execute()
    except Exception as exc:  # noqa: BLE001 - best-effort, never blocks the review response
        _log.warning("Could not insert blog review notification for %s: %s", email, exc)


@router.get("/blog/{author_type}", dependencies=[Depends(_require_admin)])
async def admin_list_blog_posts(
    author_type: Literal["artist", "tunefry"],
    status_filter: Optional[str] = Query(default=None, alias="status"),
) -> list[AdminPost]:
    try:
        rows = blog_service.list_admin_posts(author_type, status_filter)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not fetch posts: {exc}",
        ) from exc
    return [AdminPost(**r) for r in rows]


@router.post("/blog/images", dependencies=[Depends(_require_admin)])
async def admin_upload_blog_image(file: UploadFile = File(...)) -> dict[str, str]:
    key = await blog_service.save_blog_image(file)
    return {"key": key}


@router.post("/blog/rewrite", dependencies=[Depends(_require_admin)])
async def admin_rewrite_blog_post(body: RewriteRequest) -> RewriteResponse:
    new_title, new_body = await blog_ai.rewrite_article(body.title, body.body)
    return RewriteResponse(title=new_title, body=new_body)


@router.patch("/blog/{post_id}", dependencies=[Depends(_require_admin)])
async def admin_review_blog_post(post_id: str, body: AdminReviewRequest) -> AdminPost:
    try:
        row = blog_service.review_post(
            post_id,
            status_value=body.status,
            admin_note=body.admin_note,
            final_title=body.final_title,
            final_body=body.final_body,
        )
    except ReviewError as exc:
        raise _blog_review_error_to_http(exc) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not review post: {exc}",
        ) from exc

    email = (row.get("author_email") or "").lower()
    if body.status == "approved":
        _notify_artist_blog_review(
            email,
            "Your Tunefry Daily article was approved",
            f'"{row.get("final_title")}" is now live on Tunefry Daily.',
        )
    else:
        _notify_artist_blog_review(
            email,
            "Your Tunefry Daily article was declined",
            body.admin_note,
        )
    return AdminPost(**row)


@router.post("/blog/tunefry", dependencies=[Depends(_require_admin)])
async def admin_publish_tunefry_post(body: AdminTunefryPublishRequest) -> AdminPost:
    try:
        row = blog_service.create_tunefry_post(
            final_title=body.final_title,
            final_body=body.final_body,
            category=body.category,
            cover_image_keys=body.cover_image_keys,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not publish post: {exc}",
        ) from exc
    return AdminPost(**row)


@router.patch("/blog/{post_id}/flags", dependencies=[Depends(_require_admin)])
async def admin_set_blog_flags(post_id: str, body: AdminFlagsRequest) -> AdminPost:
    try:
        row = blog_service.set_flags(post_id, is_featured=body.is_featured, is_popular=body.is_popular)
    except ReviewError as exc:
        raise _blog_review_error_to_http(exc) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"Could not update flags: {exc}",
        ) from exc
    return AdminPost(**row)
